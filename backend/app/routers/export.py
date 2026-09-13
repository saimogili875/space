from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

from app.config import MOIL_MINES
from app.services.predictor import PredictionService
from app.services.anomaly import detect_anomalies

router = APIRouter(prefix="/api/export", tags=["export"])

HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CRITICAL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
ONTRACK_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _alert_fill(level):
    if level == "CRITICAL":
        return CRITICAL_FILL
    elif level == "WARNING":
        return WARNING_FILL
    return ONTRACK_FILL


@router.get("/excel")
def export_excel(mine_id: str = None):
    wb = Workbook()
    service = PredictionService()

    mines_to_export = [mine_id] if mine_id else list(MOIL_MINES.keys())

    # --- Sheet 1: Forecast ---
    ws_fc = wb.active
    ws_fc.title = "Production Forecast"
    headers_fc = ["Mine", "State", "Type", "Period", "Predicted (t)", "Target (t)", "Shortfall %", "Alert Level", "Top Factor", "Confidence Low", "Confidence High"]
    ws_fc.append(headers_fc)
    _style_header(ws_fc, len(headers_fc))

    for mid in mines_to_export:
        info = MOIL_MINES[mid]
        forecasts = service.get_forecast(mid, 4)
        for fc in forecasts:
            cause = fc.top_factors[0]["feature"] if fc.top_factors else "N/A"
            row_num = ws_fc.max_row + 1
            ws_fc.append([
                info["name"], info["state"], info["type"],
                fc.period, round(fc.predicted_tonnes), round(fc.target_tonnes),
                round(fc.shortfall_pct, 1), fc.alert_level.value, cause,
                round(fc.confidence_lower), round(fc.confidence_upper),
            ])
            for c in range(1, len(headers_fc) + 1):
                cell = ws_fc.cell(row=row_num, column=c)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")
            ws_fc.cell(row=row_num, column=8).fill = _alert_fill(fc.alert_level.value)

    _auto_width(ws_fc)

    # --- Sheet 2: Alerts ---
    ws_al = wb.create_sheet("Alerts & Actions")
    headers_al = ["Mine", "Period", "Alert Level", "Shortfall %", "Primary Cause", "Corrective Actions"]
    ws_al.append(headers_al)
    _style_header(ws_al, len(headers_al))

    for mid in mines_to_export:
        alerts = service.get_alerts(mid)
        for a in alerts:
            actions_str = "; ".join([act.action for act in a.actions[:3]]) if a.actions else "N/A"
            row_num = ws_al.max_row + 1
            ws_al.append([
                a.mine_name, a.period, a.alert_level.value,
                round(a.shortfall_pct, 1), a.primary_cause, actions_str,
            ])
            for c in range(1, len(headers_al) + 1):
                cell = ws_al.cell(row=row_num, column=c)
                cell.border = THIN_BORDER
            ws_al.cell(row=row_num, column=3).fill = _alert_fill(a.alert_level.value)

    _auto_width(ws_al)

    # --- Sheet 3: Anomaly Report ---
    ws_an = wb.create_sheet("Anomaly Report")
    headers_an = ["Mine", "Health Score", "Anomaly Count", "Top Anomaly Type", "Top Anomaly Severity", "Description"]
    ws_an.append(headers_an)
    _style_header(ws_an, len(headers_an))

    for mid in mines_to_export:
        info = MOIL_MINES[mid]
        result = detect_anomalies(mid, 12)
        health = result["summary"].get("health_score", 100)
        anomalies = result["anomalies"]
        if anomalies:
            for a in anomalies[:3]:
                row_num = ws_an.max_row + 1
                ws_an.append([
                    info["name"], health, len(anomalies),
                    a.get("type", "N/A"), a.get("severity", "N/A"), a.get("description", "N/A"),
                ])
                for c in range(1, len(headers_an) + 1):
                    ws_an.cell(row=row_num, column=c).border = THIN_BORDER
        else:
            row_num = ws_an.max_row + 1
            ws_an.append([info["name"], health, 0, "None", "None", "No anomalies detected"])
            for c in range(1, len(headers_an) + 1):
                ws_an.cell(row=row_num, column=c).border = THIN_BORDER

    _auto_width(ws_an)

    # --- Sheet 4: Mine Summary ---
    ws_ms = wb.create_sheet("Mine Summary")
    headers_ms = ["Mine", "State", "Type", "Depth (m)", "Monthly Target (t)", "Ore Grade %", "Active Since", "Health Score"]
    ws_ms.append(headers_ms)
    _style_header(ws_ms, len(headers_ms))

    for mid in mines_to_export:
        info = MOIL_MINES[mid]
        result = detect_anomalies(mid, 12)
        health = result["summary"].get("health_score", 100)
        row_num = ws_ms.max_row + 1
        ws_ms.append([
            info["name"], info["state"], info["type"], info["depth_m"],
            info["base_monthly_tonnes"], info["ore_grade_pct"], info["active_since"], health,
        ])
        for c in range(1, len(headers_ms) + 1):
            ws_ms.cell(row=row_num, column=c).border = THIN_BORDER
            ws_ms.cell(row=row_num, column=c).alignment = Alignment(horizontal="center")

    _auto_width(ws_ms)

    # Write to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"MangaLens_Report_{mine_id or 'Fleet'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
